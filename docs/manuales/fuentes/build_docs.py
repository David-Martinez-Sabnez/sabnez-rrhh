from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/Users/davidmartinez/Downloads/sabnez-rrhh')
OUT, IMG = ROOT/'docs/manuales', ROOT/'docs/manuales/capturas'
DARK, GRAY, BORDER = '1D2D3E', 'F5F6F7', 'D9D9D9'

def shade(cell, fill):
    p = cell._tc.get_or_add_tcPr(); e = OxmlElement('w:shd'); e.set(qn('w:fill'), fill); p.append(e)

def base(title, subtitle, audience):
    d=Document(); s=d.sections[0]; s.page_width=Inches(8.5); s.page_height=Inches(11)
    s.top_margin=s.bottom_margin=Inches(.72); s.left_margin=s.right_margin=Inches(.78)
    d.styles['Normal'].font.name='Aptos'; d.styles['Normal'].font.size=Pt(10.5)
    d.styles['Normal'].paragraph_format.space_after=Pt(6)
    for n,z in [('Title',25),('Heading 1',18),('Heading 2',14),('Heading 3',11.5)]:
        st=d.styles[n]; st.font.name='Aptos Display'; st.font.size=Pt(z); st.font.bold=True
        st.font.color.rgb=RGBColor(0,0,0); st.paragraph_format.space_before=Pt(12); st.paragraph_format.space_after=Pt(6); st.paragraph_format.keep_with_next=True
        ppr=st.element.get_or_add_pPr(); old=ppr.find(qn('w:pBdr'))
        if old is not None: ppr.remove(old)
    d.add_paragraph(title, style='Title'); p=d.add_paragraph(); r=p.add_run(subtitle); r.bold=True; r.font.size=Pt(13)
    p=d.add_paragraph(); p.add_run('Audiencia  ').bold=True; p.add_run(audience)
    p=d.add_paragraph(); p.add_run('Versión documentada  ').bold=True; p.add_run('Sabnez RRHH 1.3.2  |  8 de septiembre de 2026')
    d.add_page_break(); return d

def intro(d, scope, result):
    d.add_heading('Propósito y alcance',1); d.add_paragraph(scope)
    p=d.add_paragraph(); p.add_run('Resultado principal. ').bold=True; p.add_run(result)

def bullets(d, items, numbered=False):
    for x in items: d.add_paragraph(x, style='List Number' if numbered else 'List Bullet')

def table(d, headers, rows):
    t=d.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=True
    pr=t._tbl.tblPr; bd=OxmlElement('w:tblBorders'); pr.append(bd)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),BORDER); bd.append(e)
    for c,h in zip(t.rows[0].cells,headers):
        shade(c,DARK); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; r=c.paragraphs[0].add_run(h); r.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9)
    trpr=t.rows[0]._tr.get_or_add_trPr(); rep=OxmlElement('w:tblHeader'); rep.set(qn('w:val'),'true'); trpr.append(rep)
    for n,row in enumerate(rows):
        cs=t.add_row().cells
        trpr=t.rows[-1]._tr.get_or_add_trPr(); keep=OxmlElement('w:cantSplit'); trpr.append(keep)
        for c,v in zip(cs,row):
            if n%2: shade(c,GRAY)
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; c.text=str(v)
            for r in c.paragraphs[0].runs: r.font.size=Pt(8.6)
    d.add_paragraph(); return t

def figure(d, name, caption):
    pth=IMG/name
    if not pth.exists(): return
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(pth),width=Inches(6.65))
    p=d.add_paragraph(caption); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs: r.italic=True; r.font.size=Pt(8.5)

APPS=[
('Gestión de Empleados','RR. HH.','Ficha integral, contratos, documentos, bancos y dependientes.'),
('Selección Home Office','Empleado','Selección semanal de días remotos.'),('Reporte Home Office','RR. HH. y tiempos','Consulta y configuración de cupos con histórico.'),
('Mis ausencias','Empleado','Solicitudes, soportes y seguimiento.'),('Centro de aprobaciones','Aprobador','Decisiones, backups, delegaciones e historial.'),
('Estructura Comercial','Administración comercial','Clientes, proyectos, asignaciones, tarifas y aprobadores.'),
('Mis tiempos','Empleado y prestador','Registro, soportes y envío de horas.'),('Reportes de tiempos','Administración','Consulta y entregables de horas.'),
('Control Financiero','Finanzas','Devengado, facturación, recaudo, cartera, TRM y margen.'),
('Cuentas de cobro','Prestador','Generación, firma, soporte, corrección e histórico.'),
('Reporte de cuentas de cobro','RR. HH.','Expedientes y envío por lote a contabilidad.'),('Centro de configuración','Administrador','Parámetros operativos sin despliegue.')]

def manual():
    d=base('Manual de usuario de Sabnez Enterprise Platform','Guía operativa de las aplicaciones administrativas','Empleados, prestadores, aprobadores, RR. HH., administración comercial y finanzas')
    intro(d,'Explica para qué sirve cada aplicación y las tareas principales de la versión actual. Las capturas pertenecen al ambiente de validación y pueden variar en SAP Build Work Zone.','La plataforma enlaza la vinculación de personas, el registro y aprobación de tiempos, las cuentas de cobro y el control financiero.')
    d.add_heading('Cómo usar este manual',1); bullets(d,['Busque la aplicación por su nombre en SAP Build Work Zone.','Los botones dependen del estado y de los permisos.','Si una acción no aparece, revise el responsable y el estado antes de solicitar soporte.'])
    d.add_heading('Mapa de aplicaciones',1); table(d,['Aplicación','Usuario','Finalidad'],APPS)
    d.add_heading('Gestión de Empleados',1); d.add_paragraph('RR. HH. completa la ficha en borrador y la activa cuando la información está lista.')
    bullets(d,['Abra la persona y complete las pestañas de información personal, laboral, médica y de contacto.','Cree el contrato con cargo, salario y vigencia.','Cargue y clasifique anexos dentro del contrato.','Para prestadores, complete RUT, banco, dependientes y datos de cobro.','Guarde y active el borrador.'],True)
    figure(d,'empleados-datos-personales.png','Datos personales y ciudad de expedición de la identificación.')
    figure(d,'empleados-datos-cuenta-cobro.png','Datos tributarios utilizados en la cuenta de cobro.')
    figure(d,'empleados-contrato-anexos.png','Anexos asociados directamente al contrato.')
    d.add_heading('Home Office',1); bullets(d,['El empleado revisa cupos, selecciona días y confirma el estado.','RR. HH. filtra el reporte y usa Configurar cupos.','El cambio puede aplicar a una fecha o desde una fecha en adelante.','Se registran motivo, usuario, fecha y valores modificados.'],True)
    figure(d,'home-office-configuracion-cupos.png','Configuración administrativa de cupos.')
    d.add_heading('Ausencias y aprobaciones',1); bullets(d,['El empleado crea la ausencia, define fechas y adjunta soporte cuando se exige.','El aprobador abre Centro de aprobaciones y revisa detalle, documentos e historial.','Aprobar o Devolver solo aparece para el responsable autorizado.','El comentario y la decisión quedan trazados.'],True)
    figure(d,'aprobaciones-cuenta-cobro.png','Tarea de cuenta de cobro en el Centro de aprobaciones.')
    d.add_heading('Tiempos y estructura comercial',1); bullets(d,['Configure cliente, contrato, proyecto, ciclo, asignaciones, aprobadores y tarifas.','Las tarifas pueden tener vigencia histórica para periodos ya trabajados.','El empleado registra horas, descripción y soporte, revisa alertas y envía la semana.','El aprobador decide y el reporte consulta horas finales por periodo.'],True)
    figure(d,'estructura-comercial-asignaciones.png','Asignaciones y condiciones económicas del proyecto.')
    d.add_heading('Control Financiero',1); d.add_paragraph('Es una vista de control; no reemplaza la contabilidad ni la factura electrónica de Siigo.')
    bullets(d,['Seleccione el intervalo.','Revise devengado, facturación, recaudo, cartera y margen.','Atienda la alerta de horas facturables sin tarifa: esas horas devengan cero.','Complete la tarifa y actualice el tablero.'],True)
    figure(d,'control-financiero-tablero.png','Tablero y alerta de asignaciones que devengan cero.')
    d.add_heading('Cuenta de cobro del prestador',1); bullets(d,['Seleccione periodos cuyas horas ya estén aprobadas.','Combine conceptos si requiere una sola cuenta.','Genere por el valor bruto.','Adjunte afiliación en el primer cobro o PILA en los siguientes.','Firme y envíe a RR. HH.','Solicite corrección si encuentra un error.','Descargue la cuenta desde el histórico.'],True)
    figure(d,'cuentas-cobro-generar.png','Selección de periodos aprobados.')
    figure(d,'cuentas-cobro-firma.png','Firma electrónica del prestador.')
    figure(d,'cuentas-cobro-historico.png','Histórico de cuentas generadas.')
    d.add_heading('Reporte de cuentas de cobro',1); bullets(d,['Apruebe o devuelva desde Centro de aprobaciones.','Filtre los expedientes en Reporte de cuentas de cobro.','Compruebe Aprobada por RR. HH. y Lista para enviar.','Seleccione una o varias y use Enviar a contabilidad.','El lote conserva destinatario, usuario y fecha.'],True)
    figure(d,'cuentas-cobro-rrhh-envio.png','Validación que impide enviar a contabilidad un expediente sin el documento firmado.')
    d.add_heading('Centro de configuración',1); bullets(d,['Busque el parámetro.','Edite y guarde el borrador.','Use esta app para datos operativos como el correo de contabilidad.','Nunca guarde secretos o contraseñas.'],True)
    figure(d,'configuracion-parametros.png','Correo destinatario de contabilidad.')
    d.add_heading('Solución rápida de problemas',1); table(d,['Situación','Revisión'],[
    ('No aparece una acción','Estado, responsable actual y rol.'),('No aparece un periodo','Horas aprobadas, contrato, tarifa y banco.'),('Expediente incompleto','PDF firmado, soporte y análisis del archivo.'),('No se envía','Aprobación, correo configurado y Microsoft Graph.'),('Hora en cero','Tarifa vigente.'),('App no visible','Colección BTP y contenido de Work Zone.')])
    d.save(OUT/'Manual_de_usuario_Sabnez_Enterprise_Platform.docx')

def functional():
    d=base('Documentación funcional de Sabnez Enterprise Platform','Procesos, reglas y alcance de negocio','Dueños de proceso, RR. HH., administración comercial, finanzas y producto')
    intro(d,'Describe el comportamiento implementado y las relaciones entre datos maestros, solicitudes, aprobaciones, tiempos, cobro y finanzas.','La integración con Siigo y el cálculo automático de retenciones quedan para una fase posterior; hoy la cuenta se firma por valor bruto.')
    d.add_heading('Alcance por dominio',1); table(d,['Dominio','Capacidades'],[
    ('Personas','Ficha integral, contratos, anexos, bancos, dependientes, afiliaciones y beneficios.'),('Aprobaciones','Bandeja, asignación, backup, delegación, decisión e historial.'),('Tiempos','Estructura comercial, registro, aprobación, tarifas y reportes.'),('Home Office','Selección, cupos, políticas e histórico.'),('Finanzas','Devengado, factura, recaudo, cartera, TRM, retenciones y margen.'),('Cuentas de cobro','Elegibilidad, firma, seguridad social, corrección, aprobación y envío.'),('Configuración','Parámetros operativos editables.')])
    d.add_heading('Ciclo de cuentas de cobro',1); bullets(d,['El periodo mensual inicia el día de comienzo de la asignación y termina el día anterior del mes siguiente.','Todas las horas deben completar la aprobación exigida por el proyecto.','Contrato, tarifa y cuenta bancaria deben cubrir el periodo.','El prestador decide combinar o separar conceptos.','La cuenta usa el valor bruto; las retenciones se determinan al pagar.','El primer cobro exige afiliación y los siguientes PILA.','La firma conserva nombre, usuario, fecha y huella.','RR. HH. aprueba o devuelve desde el Centro de aprobaciones.','Las aprobadas se envían por lote a contabilidad sin reenvío accidental.'],True)
    d.add_heading('Estados de la cuenta',2); table(d,['Estado','Significado'],[('PENDING_SIGNATURE','Pendiente de firma'),('SIGNED','Firma registrada'),('SUBMITTED','Enviada a RR. HH.'),('UNDER_HR_REVIEW','En revisión'),('CORRECTION_REQUESTED','Corrección solicitada'),('HR_APPROVED','Aprobada para pago'),('SENT_TO_ACCOUNTING','Enviada a contabilidad'),('HR_REJECTED','Rechazada'),('CANCELLED','Cancelada')])
    d.add_heading('Reglas transversales',1); bullets(d,['Las acciones se calculan en el servidor y no por la interfaz.','Solo se muestran asignaciones vigentes para la fecha.','Horas enviadas se bloquean hasta ser devueltas.','Tarifas históricas permiten valorar trabajo ya ejecutado.','Cupos de Home Office no pueden quedar por debajo de la ocupación confirmada.','Los cambios y decisiones registran quién, cuándo y por qué.'])
    d.add_heading('Retenciones e integración',1); d.add_paragraph('La contadora calcula retenciones en Siigo al pagar y envía el comprobante al prestador. Los datos tributarios preparan esa integración, pero la plataforma actual no descuenta retenciones de la cuenta firmada.')
    d.add_heading('Pendientes funcionales',1); table(d,['Tema','Estado','Evolución'],[('Siigo','No implementado','API, conciliación, errores y comprobantes.'),('Retenciones','Diferidas al pago','Implementar junto con Siigo.'),('Work Zone','Notificaciones deshabilitadas','Provisionar SAP_Notifications.'),('Capturas restantes','Pendientes','Completar en Work Zone con datos representativos.'),('Validación real','En progreso','Ejecutar un ciclo por proceso y rol.')])
    d.save(OUT/'Documentacion_funcional_Sabnez_Enterprise_Platform.docx')

def technical():
    d=base('Documentación técnica de Sabnez Enterprise Platform','Arquitectura, componentes, despliegue y operación','Desarrollo, administración SAP BTP y soporte')
    intro(d,'Documenta el proyecto sabnez-rrhh, sus módulos CAP, aplicaciones UI5, persistencia, integraciones y operación.','Es una aplicación MTA: CAP Node.js expone OData V4, PostgreSQL persiste, XSUAA autoriza y HTML5 Repository distribuye doce aplicaciones.')
    d.add_heading('Arquitectura',1); bullets(d,['SAPUI5 y Fiori Elements en SAP Build Work Zone.','SAP CAP Node.js y servicios OData V4.','SQLite local y PostgreSQL productivo.','XSUAA en producción; autenticación simulada en desarrollo y pruebas.','Adjuntos con análisis antimalware.','Microsoft Graph para correo; Job Scheduler para recordatorios.','Notificaciones Work Zone preparadas pero temporalmente deshabilitadas.'])
    d.add_heading('Aplicaciones y servicios',1); table(d,['Aplicación','Servicio','Ruta'],[
    ('empleadosui','AdminService','/admin'),('homeofficeui','HomeOfficeService','/home-office'),('homeofficeadminui','HomeOfficeAdminService','/home-office-admin'),('ausenciasui','EmployeeAbsenceService','/ausencias-empleado'),('aprobacionesui','ApprovalService','/aprobaciones'),('tiemposadminui','TimeAdminService','/tiempos-admin'),('tiemposempleadoui','EmployeeTimeService','/tiempos-empleado'),('reportestiemposui','TimeReportingService','/time-reports'),('finanzasui','FinanceService','/finanzas'),('cuentascobroui','CollectionAccountService','/cuentas-cobro'),('cuentascobrorrhhui','CollectionAccountService','/cuentas-cobro'),('configuracionui','ConfiguracionService','/configuracion')])
    d.add_heading('Modelo de datos',1); table(d,['Área','Entidades principales'],[('RR. HH.','Empleados, contratos, bancos, dependientes, documentos, contactos, ausencias y saldos.'),('Aprobaciones','Definiciones, etapas, instancias, tareas, asignaciones, decisiones, delegaciones y eventos.'),('Tiempos','Clientes, proyectos, reglas, ciclos, asignaciones, tarifas, hojas, registros y revisiones.'),('Cobro','CollectionAccounts, Items y Events con instantáneas históricas.'),('Finanzas','Facturas, líneas, retenciones, recaudos, tasas y compensatorios.'),('Home Office','Configuración, políticas, historial, semanas, días y selecciones.'),('Configuración','Parametros por clave, grupo, tipo y valor.')])
    d.add_heading('Repositorio',1); table(d,['Ruta','Contenido'],[('app/','Aplicaciones UI5 y anotaciones.'),('db/','Modelo CDS y semillas.'),('srv/','Servicios y lógica de dominio.'),('test/','21 pruebas unitarias e integración.'),('scripts/','Migraciones, sincronización y limpieza.'),('docs/','Decisiones y manuales.'),('mta.yaml','Módulos, recursos y colecciones.'),('xs-security.json','Scopes y roles.')])
    d.add_heading('Construcción y despliegue',1); bullets(d,['Ejecutar pruebas y compilación CDS.','Limpiar artefactos.','Construir MTA.','Desplegar en Cloud Foundry.','Verificar deployer PostgreSQL y una instancia activa del servicio.','Publicar Work Zone y asignar colecciones.','Probar con usuarios representativos.'],True)
    table(d,['Objetivo','Comando'],[('Local','npm run watch'),('Pruebas','npm test'),('Compilar','npm run compile'),('Limpiar','npm run clean:build'),('Construir','npm run build'),('Desplegar','npm run deploy'),('Sincronizar SQLite','npm run migrate:sqlite'),('Recrear SQLite','npm run reset:sqlite')])
    d.add_heading('Configuración e integraciones',1); bullets(d,['Parámetros funcionales como CONTABILIDAD_EMAIL viven en Centro de configuración.','GRAPH_TENANT_ID, GRAPH_CLIENT_ID y GRAPH_CLIENT_SECRET permanecen en Cloud Foundry.','Los destinos ui5 y srv-api distribuyen UI5 y reenvían el token al backend.','Credenciales locales se excluyen de Git.'])
    d.add_heading('Pruebas y operación',1); d.add_paragraph('Los 21 archivos de test validan ausencias, aprobaciones, facturación, cobro, clasificación, compensatorios, matemáticas financieras, correo, TRM, streams y tiempos. No se usan en producción, pero deben conservarse.')
    bullets(d,['Monitorear correo, análisis de archivos, tareas y recordatorios.','Agregar pruebas negativas por rol, concurrencia e idempotencia.','Probar restauración de PostgreSQL y recuperación de documentos.'])
    d.save(OUT/'Documentacion_tecnica_Sabnez_Enterprise_Platform.docx')

def security():
    d=base('Seguridad y roles de Sabnez Enterprise Platform','Modelo de acceso para SAP BTP y controles operativos','Administración BTP, seguridad, RR. HH., finanzas y soporte')
    intro(d,'Define scopes, colecciones, asignación recomendada y controles de XSUAA, CAP, archivos, correo y operación.','La autorización se aplica en CAP; ocultar una aplicación o un botón no sustituye el control del backend.')
    d.add_heading('Principios',1); bullets(d,['Mínimo privilegio y separación de funciones.','Autogestión limitada a la identidad autenticada.','Autorización en cada lectura, descarga, decisión y envío.','Trazabilidad de cambios y decisiones.','Secretos fuera del repositorio y del Centro de configuración.'])
    d.add_heading('Roles de contenido en SAP Build Work Zone',1)
    d.add_paragraph('Estos roles organizan la visibilidad de páginas, catálogos, aplicaciones y tarjetas en Work Zone. Deben combinarse con las colecciones BTP correspondientes: asignar un contenido en Work Zone no concede por sí solo autorización sobre los servicios CAP.')
    table(d,['Rol de Work Zone','ID','Contenido esperado','Colección BTP relacionada'],[
    ('Administrador de Configuraciones','Administrador_de_Configuraciones','Centro de configuración.','RRHH Admin.'),
    ('Prestadores de servicios','Prestadores_de_servicios','Zona Colaboradores y generación de cuentas de cobro.','Usuario autenticado; CollectionAccountHR no se asigna al prestador.'),
    ('Finanzas','Finanzas','Control financiero, tarifas, costos y rentabilidad.','Tiempos Finanzas.'),
    ('General','~sbz.wz.cards_sbz.wz.cards.general.role','Tarjetas generales del canal Sabnez Cards.','Depende de la acción de cada tarjeta.'),
    ('Everyone','sap_subaccount_everyone','Contenido común no sensible para todos los usuarios.','Ninguna colección administrativa.'),
    ('Administrativos RRHH','Administrativos_RRHH','Gestión de empleados, documentos y revisión de cuentas.','RRHH Editor y Cuentas de cobro RRHH; RRHH Admin solo cuando corresponda.'),
    ('Empleados Sabnez','Empleados_Sabnez','Autogestión del empleado, tiempos, ausencias y Home Office.','Usuario autenticado y filtrado por identidad.'),
    ('Gestión de proyectos','Gestion_proyectos','Estructura comercial, proyectos, aprobaciones y reportes operativos.','Tiempos Admin; agregar Tiempos Finanzas únicamente si consulta valores económicos.')])
    figure(d,'roles-work-zone.png','Roles de contenido configurados en SAP Build Work Zone.')
    d.add_heading('Colecciones BTP',1); table(d,['Colección','Plantilla','Uso'],[('RRHH Editor','Editor','Crear y modificar empleados y documentos.'),('RRHH Admin','Admin','Eliminar y administrar parámetros.'),('Aprobaciones Admin','ApprovalAdmin','Delegaciones y operación de aprobaciones.'),('Tiempos Admin','TimeAdmin','Estructura, tiempos, Home Office y reportes.'),('Tiempos Finanzas','TimeFinance','Tarifas, costos, margen y finanzas.'),('Entregables Tiempos','TimeDeliverables','Generar entregables.'),('Cuentas de cobro RRHH','CollectionAccountHR','Expedientes y envío a contabilidad.')])
    d.add_heading('Roles adicionales',2); table(d,['Plantilla','Asignación'],[('TimeApprovalReassign','Solo soporte. La plantilla existe, pero falta una colección predefinida en el MTA.'),('CollectionAccountJob','Solo Job Scheduler; nunca usuarios humanos.'),('authenticated-user','Autogestión filtrada por identidad.')])
    d.add_heading('Perfiles recomendados',1); table(d,['Perfil','Colecciones mínimas'],[('Empleado','Ninguna administrativa.'),('Prestador','Ninguna administrativa.'),('RR. HH. operativo','RRHH Editor + Cuentas de cobro RRHH.'),('RR. HH. administrador','RRHH Editor + RRHH Admin + Cuentas de cobro RRHH.'),('Comercial','Tiempos Admin.'),('Finanzas','Tiempos Finanzas; Tiempos Admin solo si administra estructura.'),('Aprobador','Autenticado; la tarea vigente define su decisión.'),('Administrador de aprobaciones','Aprobaciones Admin.'),('Soporte','TimeApprovalReassign temporal.')])
    d.add_heading('Acceso por servicio',1); table(d,['Servicio','Requisito'],[('AdminService','Editor; DELETE requiere Admin.'),('ConfiguracionService','Admin.'),('HomeOfficeAdminService','Editor y TimeAdmin.'),('TimeAdminService','TimeAdmin; economía requiere TimeFinance.'),('FinanceService','TimeFinance.'),('TimeReportingService','TimeAdmin; entregables requieren TimeDeliverables.'),('CollectionAccountService','Autenticado; RR. HH. requiere CollectionAccountHR.'),('Aprobaciones','Autenticado; reasignación requiere TimeApprovalReassign.'),('Autogestión','Autenticado y filtrado por identidad.')])
    d.add_heading('Archivos y datos sensibles',1); bullets(d,['Contratos, certificados médicos, liquidaciones, PILA y cuentas firmadas son confidenciales.','La descarga debe validar empleado, proyecto o rol.','No habilitar archivos pendientes o rechazados por antimalware.','Validar tamaño, tipo real, nombre seguro y contenido.','Evitar salarios, secretos y documentos completos en logs y correos.'])
    d.add_heading('Microsoft Graph',1); bullets(d,['Tenant, client y secret son credenciales técnicas.','El correo de contabilidad es parámetro funcional, no secreto.','Conceder solo permiso de envío desde el buzón autorizado.','Rotar el secreto y vigilar su vencimiento.'])
    d.add_heading('Pendientes prioritarios',1); table(d,['Prioridad','Acción'],[('Alta','Crear colección para TimeApprovalReassign y limitarla a soporte.'),('Alta','Completar pruebas negativas de acceso cruzado.'),('Alta','Definir conservación de soportes médicos, PILA y contratos.'),('Media','Provisionar SAP_Notifications antes de habilitar Work Zone.'),('Media','Formalizar rotación de Microsoft Graph.'),('Media','Revisar trimestralmente las colecciones asignadas.')])
    d.save(OUT/'Seguridad_y_roles_SAP_BTP_Sabnez_Enterprise_Platform.docx')

if __name__=='__main__': OUT.mkdir(parents=True,exist_ok=True); manual(); functional(); technical(); security()
